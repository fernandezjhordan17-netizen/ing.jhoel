"""Exportación de resultados a Excel (memoria de cálculo, entregables)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Font


def espectro_a_excel(espectro: dict, ruta: str) -> str:
    """Hoja 'Parametros' + hoja 'Espectro' con gráfico Sa/g vs T."""
    destino = Path(ruta).expanduser()
    if destino.suffix.lower() != ".xlsx":
        destino = destino.with_suffix(".xlsx")
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    hp = wb.active
    hp.title = "Parametros"
    hp.append(["Norma", espectro["norma"]])
    for k, v in espectro["parametros"].items():
        hp.append([k, v])
    hp.append([])
    hp.append(["Notas"])
    for n in espectro["notas"] + espectro.get("pendientes", []):
        hp.append([n])
    hp["A1"].font = Font(bold=True)
    he = wb.create_sheet("Espectro")
    he.append(["T (s)", "Sa/g horizontal", "Sa/g vertical"])
    for t, sa, sv in zip(espectro["periodos_s"], espectro["sa_g"], espectro["sa_vertical_g"]):
        he.append([t, sa, sv])
    for c in he[1]:
        c.font = Font(bold=True)
    graf = ScatterChart()
    graf.title = "Espectro de diseño E.030-2026"
    graf.x_axis.title = "T (s)"
    graf.y_axis.title = "Sa/g"
    n = len(espectro["periodos_s"]) + 1
    serie = Series(Reference(he, min_col=2, min_row=1, max_row=n), Reference(he, min_col=1, min_row=2, max_row=n),
                   title_from_data=True)
    graf.series.append(serie)
    he.add_chart(graf, "E2")
    wb.save(destino)
    return str(destino)
