"""
Metrados (cantidades) desde modelos IFC usando los conjuntos de cantidades (Qto) del modelo,
con unidades convertidas a SI y partidas sugeridas según la Norma Técnica de Metrados (RD 073-2010-VIVIENDA).
El encofrado no viene en los Qto estándar: se reporta como pendiente de cálculo.
"""
from __future__ import annotations

import math
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from .ifc import abrir, _ifcopenshell

DENSIDAD_ACERO = 7850.0  # kg/m³

# clase IFC -> (partida sugerida, magnitud principal)
PARTIDAS = {
    "IfcFooting": ("Estructuras – Concreto en cimentaciones (zapatas/cimientos)", "volumen"),
    "IfcPile": ("Estructuras – Pilotes", "volumen"),
    "IfcColumn": ("Estructuras – Concreto en columnas", "volumen"),
    "IfcBeam": ("Estructuras – Concreto en vigas", "volumen"),
    "IfcSlab": ("Estructuras – Concreto en losas", "volumen"),
    "IfcStair": ("Estructuras – Concreto en escaleras", "volumen"),
    "IfcStairFlight": ("Estructuras – Concreto en escaleras", "volumen"),
    "IfcMember": ("Estructuras – Elementos secundarios", "volumen"),
    "IfcPlate": ("Estructuras – Planchas", "area"),
    "IfcReinforcingBar": ("Estructuras – Acero de refuerzo fy = 420 MPa", "peso"),
    "IfcReinforcingMesh": ("Estructuras – Malla de refuerzo", "peso"),
    "IfcWall": ("Arquitectura – Muros y tabiques", "area"),
    "IfcWallStandardCase": ("Arquitectura – Muros y tabiques", "area"),
    "IfcCurtainWall": ("Arquitectura – Muro cortina", "area"),
    "IfcCovering": ("Arquitectura – Revestimientos, pisos y cielos rasos", "area"),
    "IfcRoof": ("Arquitectura – Coberturas", "area"),
    "IfcDoor": ("Arquitectura – Carpintería: puertas", "conteo"),
    "IfcWindow": ("Arquitectura – Carpintería: ventanas", "conteo"),
    "IfcRailing": ("Arquitectura – Barandas", "longitud"),
    "IfcPipeSegment": ("Instalaciones sanitarias – Tuberías", "longitud"),
    "IfcFlowSegment": ("Instalaciones – Tuberías/ductos", "longitud"),
    "IfcDuctSegment": ("Instalaciones mecánicas – Ductos", "longitud"),
    "IfcCableCarrierSegment": ("Instalaciones eléctricas – Bandejas/canaletas", "longitud"),
    "IfcSanitaryTerminal": ("Instalaciones sanitarias – Aparatos sanitarios", "conteo"),
    "IfcLightFixture": ("Instalaciones eléctricas – Luminarias", "conteo"),
    "IfcOutlet": ("Instalaciones eléctricas – Salidas/tomacorrientes", "conteo"),
    "IfcFlowTerminal": ("Instalaciones – Terminales", "conteo"),
}
UNIDADES = {"volumen": "m³", "area": "m²", "longitud": "m", "peso": "kg", "conteo": "und"}
PREFERENCIA = {
    "volumen": ("NetVolume", "GrossVolume", "Volume"),
    "area": ("NetSideArea", "NetArea", "NetSurfaceArea", "GrossSideArea", "GrossArea", "Area"),
    "longitud": ("Length", "NetLength", "GrossLength"),
    "peso": ("NetWeight", "GrossWeight", "Weight"),
}
ATRIBUTO = {"IfcQuantityVolume": ("VolumeValue", "VOLUMEUNIT"), "IfcQuantityArea": ("AreaValue", "AREAUNIT"),
            "IfcQuantityLength": ("LengthValue", "LENGTHUNIT"), "IfcQuantityWeight": ("WeightValue", "MASSUNIT"),
            "IfcQuantityCount": ("CountValue", None)}


def _escalas(modelo) -> dict:
    uu = _ifcopenshell().util.unit
    esc = {}
    for t in ("LENGTHUNIT", "AREAUNIT", "VOLUMEUNIT"):
        try:
            esc[t] = uu.calculate_unit_scale(modelo, t)
        except Exception:
            esc[t] = 1.0
    try:
        esc["MASSUNIT"] = uu.calculate_unit_scale(modelo, "MASSUNIT") / 1000.0  # IfcOpenShell usa gramos como base
    except Exception:
        esc["MASSUNIT"] = 1.0
    return esc


def _cantidades(elemento, esc: dict) -> tuple[dict, int]:
    """Cantidades Qto del elemento en SI (m, m², m³, kg) y número de duplicados corregidos.

    Algunos exportadores (p. ej. Revit con proyecto en mm) escriben la misma cantidad dos veces, una en mm²/mm³
    y otra en m²/m³. Si dos valores del mismo nombre difieren en ~10^6 o ~10^9 se toma el menor (el coherente con
    la unidad declarada) y se reporta."""
    valores: dict[str, list[float]] = {}
    for rel in getattr(elemento, "IsDefinedBy", None) or []:
        if not rel.is_a("IfcRelDefinesByProperties"):
            continue
        qto = rel.RelatingPropertyDefinition
        if not qto.is_a("IfcElementQuantity"):
            continue
        for q in qto.Quantities or []:
            info = ATRIBUTO.get(q.is_a())
            if not info:
                continue
            valor = getattr(q, info[0], None)
            if valor is not None:
                valores.setdefault(q.Name, []).append(valor * (esc[info[1]] if info[1] else 1.0))
    salida, corregidos = {}, 0
    for nombre, lista in valores.items():
        positivos = [v for v in lista if v > 0]
        if len(positivos) > 1:
            razon = max(positivos) / min(positivos)
            if any(abs(razon / f - 1) < 0.02 for f in (1e6, 1e9)):
                salida[nombre] = min(positivos)
                corregidos += 1
                continue
        salida[nombre] = lista[0]
    # Plausibilidad: un elemento no supera 1e5 m² ni 1e5 m³; si Largo×Alto existe, el área debe parecerse.
    ref_area = salida["Length"] * salida["Height"] if salida.get("Length") and salida.get("Height") else None
    for nombre, valor in list(salida.items()):
        tipo = next((t for t, claves in (("area", PREFERENCIA["area"]), ("volumen", PREFERENCIA["volumen"]))
                     if nombre in claves), None)
        if tipo is None or valor <= 0:
            continue
        factor = 1e6 if tipo == "area" else 1e9
        fuera = valor > 1e5 or (tipo == "area" and ref_area and valor / ref_area > 1e4)
        if fuera and valor / factor < 1e5:
            salida[nombre] = valor / factor
            corregidos += 1
    return salida, corregidos


def _peso_barra(barra, esc: dict) -> float | None:
    d = getattr(barra, "NominalDiameter", None)
    largo = getattr(barra, "BarLength", None)
    if not d or not largo:
        return None
    d_m, l_m = d * esc["LENGTHUNIT"], largo * esc["LENGTHUNIT"]
    return math.pi / 4 * d_m ** 2 * l_m * DENSIDAD_ACERO


def metrar(ruta: str, por_piso: bool = False) -> dict:
    modelo = abrir(ruta)
    ue = _ifcopenshell().util.element
    esc = _escalas(modelo)
    grupos: dict[tuple, dict] = defaultdict(lambda: {"cantidad": 0.0, "elementos": 0, "sin_dato": 0})
    duplicados = 0
    for clase, (partida, magnitud) in PARTIDAS.items():
        try:
            elementos = modelo.by_type(clase, include_subtypes=False)
        except RuntimeError:  # clase inexistente en este esquema
            continue
        for e in elementos:
            piso = ""
            if por_piso:
                c = ue.get_container(e)
                piso = c.Name if c else "(sin piso)"
            g = grupos[(clase, partida, magnitud, piso)]
            g["elementos"] += 1
            if magnitud == "conteo":
                g["cantidad"] += 1
                continue
            qs, corregidos = _cantidades(e, esc)
            duplicados += corregidos
            valor = next((qs[n] for n in PREFERENCIA[magnitud] if n in qs), None)
            if valor is None and magnitud == "peso" and e.is_a("IfcReinforcingBar"):
                valor = _peso_barra(e, esc)
            if valor is None:
                g["sin_dato"] += 1
            else:
                g["cantidad"] += valor
    filas = [{"clase": clase, "partida": partida, "piso": piso or None, "unidad": UNIDADES[mag],
              "cantidad": round(g["cantidad"], 3), "elementos": g["elementos"], "sin_dato": g["sin_dato"]}
             for (clase, partida, mag, piso), g in sorted(grupos.items())]
    avisos = [f"{f['clase']}: {f['sin_dato']} de {f['elementos']} elementos sin cantidades Qto (exporta el IFC "
              "con 'cantidades base' o calcula desde la geometría)." for f in filas if f["sin_dato"]]
    if duplicados:
        avisos.append(f"⚠️ {duplicados} cantidades con error de unidades (mm²/mm³ declaradas como m²/m³, a veces duplicadas) "
                      "en el IFC: se corrigieron al valor coherente. Revisa la configuración del exportador IFC.")
    avisos.append("Encofrado (m²) no está en los Qto estándar: calcúlalo por caras de contacto (Dynamo/Revit) según la Norma Técnica de Metrados.")
    return {"archivo": str(Path(ruta)), "esquema": modelo.schema, "unidades_origen": {k: v for k, v in esc.items()},
            "filas": filas, "avisos": avisos, "referencia": "Norma Técnica de Metrados (RD 073-2010-VIVIENDA/VMCS/DNC)"}


def metrado_a_excel(metrado: dict, ruta: str) -> str:
    destino = Path(ruta).expanduser()
    if destino.suffix.lower() != ".xlsx":
        destino = destino.with_suffix(".xlsx")
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Metrado"
    cab = ["Partida sugerida", "Clase IFC", "Piso", "Unidad", "Cantidad", "N.° elementos", "Sin dato"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True)
    for f in metrado["filas"]:
        ws.append([f["partida"], f["clase"], f["piso"], f["unidad"], f["cantidad"], f["elementos"], f["sin_dato"]])
    ws.column_dimensions["A"].width = 60
    av = wb.create_sheet("Avisos")
    av.append(["Fuente", metrado["archivo"]])
    av.append(["Referencia", metrado["referencia"]])
    for a in metrado["avisos"]:
        av.append([a])
    wb.save(destino)
    return str(destino)
