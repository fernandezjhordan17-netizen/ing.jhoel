"""
Ejecución de un proyecto estructural de punta a punta:
E.020 (pesos por nivel) → E.030-2026 (sitio, espectro, cortante estática por dirección) → E.060 (combinaciones,
vigas, columnas) → metrados IFC (opcional) → memoria en Excel + notas del proyecto en la bóveda + bitácora.

Por defecto es un *dry-run*: calcula todo y devuelve qué archivos escribiría; solo escribe con confirmar=True.
Los datos del proyecto (JSON) son datos, no instrucciones.
"""
from __future__ import annotations

import datetime as dt
import json
import re
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill

from . import bitacora, e020, e030, e060
from .boveda import ruta_boveda

CARPETA_PROYECTOS = Path("70 PROYECTOS") / "Proyectos activos"
CARPETA_SALIDAS = Path("90 RECURSOS") / "Adjuntos" / "Proyectos"

# tipo de CT por defecto según el sistema (art. 36); None = exige el periodo del análisis modal
CT_POR_SISTEMA = {
    "concreto_porticos": "porticos", "acero_smf": "porticos", "acero_imf": "porticos", "acero_omf": "porticos",
    "acero_scbf": "acero_arriostrado", "acero_ocbf": "acero_arriostrado", "acero_ebf": "acero_arriostrado",
    "concreto_dual": "dual", "concreto_muros": "muros", "concreto_emdl": "emdl", "albanileria": "albanileria",
    "madera": None, "pendulo_invertido": None,
}
MUROS_PORTANTES = {"concreto_muros", "concreto_emdl", "albanileria"}
MATERIAL_DERIVA = {"concreto_emdl": "emdl", "albanileria": "albanileria", "madera": "madera"}


class ErrorProyecto(ValueError):
    """Datos del proyecto incompletos o incoherentes."""


# ------------------------------------------------------------------ utilidades -----------------
def nombre_archivo(texto: str) -> str:
    """Sin tildes ni caracteres prohibidos en Obsidian/Windows (convención de la bóveda)."""
    plano = "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")
    plano = re.sub(r'[:*?"<>|/\\#^\[\]]', "-", plano)
    return re.sub(r"\s+", " ", plano).strip(" .-")


def _num(x: float, d: int = 2) -> str:
    return f"{x:,.{d}f}".replace(",", " ").replace(".", ",")


def _g(x) -> str:
    """Número con coma decimal y sin ceros sobrantes (texto de la bóveda)."""
    if isinstance(x, bool) or not isinstance(x, (int, float)):
        return "" if x is None else str(x)
    return f"{x:.5f}".rstrip("0").rstrip(".").replace(".", ",")


def cargar(ruta: str | Path) -> dict:
    ruta = Path(ruta).expanduser()
    datos = json.loads(ruta.read_text(encoding="utf-8"))
    datos.setdefault("_origen", str(ruta))
    return datos


def _requerir(cfg: dict, *claves: str) -> None:
    faltan = [c for c in claves if cfg.get(c) in (None, "", [])]
    if faltan:
        raise ErrorProyecto(f"Faltan datos del proyecto: {', '.join(faltan)}.")


# ------------------------------------------------------------------ E.020: pesos -------------------
def pesos_por_nivel(cfg: dict) -> dict:
    """Carga muerta, viva y peso sísmico por nivel (art. 31 E.030; cargas E.020)."""
    _requerir(cfg, "niveles")
    sismo, base = cfg["sismo"], cfg.get("cargas", {})
    categoria = sismo["categoria"].upper()
    filas, altura = [], 0.0
    for i, nv in enumerate(cfg["niveles"]):
        c = {**base, **nv}
        if "altura_entrepiso_m" not in c:
            raise ErrorProyecto(f"Nivel {i + 1}: falta altura_entrepiso_m.")
        altura += c["altura_entrepiso_m"]
        azotea = bool(c.get("azotea"))
        fuentes = {}
        if "cm_kn" in c:  # peso directo (p. ej. masa del modelo ETABS)
            cm, cv = float(c["cm_kn"]), float(c.get("cv_kn", 0.0))
            area, cm_kpa, cv_kpa = c.get("area_m2"), None, None
            fuentes["CM"] = "dato del proyecto (cm_kn)"
        else:
            if "area_m2" not in c:
                raise ErrorProyecto(f"Nivel {i + 1}: indica area_m2 (o cm_kn/cv_kn directos).")
            area = float(c["area_m2"])
            componentes = {}
            if c.get("aligerado_m"):
                componentes["aligerado"] = e020.peso_aligerado(float(c["aligerado_m"]))
                fuentes["aligerado"] = e020.REFERENCIAS["aligerado"]
            for k in ("losa_kpa", "acabados_kpa", "tabiqueria_kpa", "estructura_kpa", "otros_cm_kpa"):
                if c.get(k):
                    componentes[k.removesuffix("_kpa")] = float(c[k])
                    fuentes[k.removesuffix("_kpa")] = "supuesto del proyecto (revisar)"
            cm_kpa = sum(componentes.values())
            if "cv_kpa" in c:
                cv_kpa, fuentes["CV"] = float(c["cv_kpa"]), "dato del proyecto"
            elif azotea:
                cv_kpa, fuentes["CV"] = e020.CARGA_VIVA_TECHO, e020.REFERENCIAS["techo"]
            else:
                cv_kpa = e020.carga_viva(c.get("uso", "vivienda"))
                fuentes["CV"] = f"{e020.REFERENCIAS['carga_viva']} ({c.get('uso', 'vivienda')})"
            cm, cv = cm_kpa * area, cv_kpa * area
        fraccion = 0.25 if azotea else e030.FRACCION_CV[categoria]
        filas.append({
            "nivel": c.get("nombre", f"Nivel {i + 1}"), "altura_m": round(altura, 3),
            "entrepiso_m": c["altura_entrepiso_m"], "area_m2": area, "azotea": azotea,
            "cm_kpa": None if cm_kpa is None else round(cm_kpa, 3),
            "cv_kpa": None if cv_kpa is None else round(cv_kpa, 3),
            "CM_kN": round(cm, 2), "CV_kN": round(cv, 2), "fraccion_CV": fraccion,
            "P_kN": round(cm + fraccion * cv, 2), "fuentes": fuentes,
        })
    total = sum(f["P_kN"] for f in filas)
    return {"filas": filas, "P_total_kN": round(total, 2), "hn_m": round(altura, 3), "pisos": len(filas),
            "criterio": f"P = CM + {int(100 * e030.FRACCION_CV[categoria])} % CV (categoría {categoria}); "
                        "azotea 25 % CV (E.030 art. 31)"}


# ------------------------------------------------------------------ E.030 por dirección ---------------
def analisis_sismico(cfg: dict, pesos: dict) -> dict:
    s = cfg["sismo"]
    _requerir(s, "zona", "categoria")
    zona, cat = int(s["zona"]), s["categoria"].upper()
    comunes = dict(perfil=s.get("perfil"), vs30=s.get("vs30"), ts=s.get("ts"), aislamiento=bool(s.get("aislamiento")))
    direcciones = s.get("direcciones") or {"X": {}, "Y": {}}
    resultados = {}
    for d, dd in direcciones.items():
        sistema = dd.get("sistema", s.get("sistema"))
        if not sistema:
            raise ErrorProyecto(f"Dirección {d}: indica el sistema estructural.")
        irr = dd.get("irregularidades", s.get("irregularidades", [])) or []
        regular = not irr
        if dd.get("periodo_s"):
            T, fuente_t = float(dd["periodo_s"]), "análisis modal (dato del proyecto)"
        else:
            tipo_ct = dd.get("tipo_ct") or CT_POR_SISTEMA.get(sistema)
            if tipo_ct is None:
                raise ErrorProyecto(f"Dirección {d}: el sistema {sistema} no tiene CT; indica periodo_s del modal.")
            T = e030.periodo_aproximado(pesos["hn_m"], tipo_ct)
            fuente_t = f"T = hn/CT = {pesos['hn_m']}/{e030.CT[tipo_ct]:g} (art. 36)"
        estatica = e030.cortante_estatica(
            zona, cat, sistema, pesos["P_total_kN"], T, irregularidades=irr, **comunes,
            pesos_niveles=[f["P_kN"] for f in pesos["filas"]], alturas_niveles=[f["altura_m"] for f in pesos["filas"]])
        esp = e030.espectro(zona, cat, sistema, irregularidades=irr, **comunes)
        restricciones = e030.verificar_restricciones(cat, zona, sistema, irr, pisos=pesos["pisos"],
                                                     altura_m=pesos["hn_m"], perfil=esp["parametros"]["perfil"],
                                                     aislamiento=comunes["aislamiento"])
        hn = pesos["hn_m"]
        estatico_ok = zona == 1 or (regular and hn <= 30) or (sistema in MUROS_PORTANTES and hn <= 15)
        r = {
            "sistema": sistema, "descripcion": e030.DESCRIPCION_SISTEMAS[sistema], "irregularidades": irr,
            "regular": regular, "T_s": round(T, 4), "fuente_T": fuente_t, "estatica": estatica, "espectro": esp,
            "restricciones": restricciones, "metodo_estatico_aplicable": estatico_ok,
            "V_minima_dinamica_kN": round((0.80 if regular else 0.90) * estatica["V"], 2),
        }
        res_etabs = (cfg.get("resultados_etabs") or {}).get(d, {})
        if res_etabs.get("v_dinamica_kN"):
            r["escalamiento"] = e030.escalamiento_dinamico(float(res_etabs["v_dinamica_kN"]), estatica["V"], regular)
        if res_etabs.get("derivas_elasticas"):
            material = res_etabs.get("material") or MATERIAL_DERIVA.get(sistema, "concreto")
            r["derivas"] = e030.verificar_derivas(res_etabs["derivas_elasticas"], estatica["R"], material, regular,
                                                  res_etabs.get("etiquetas"))
        resultados[d] = r
    return resultados


# ------------------------------------------------------------------ E.060 ---------------------------
def diseno_concreto(cfg: dict) -> dict:
    mat = cfg.get("materiales", {})
    fc, fy = float(mat.get("fc_mpa", 21)), float(mat.get("fy_mpa", 420))
    casos = cfg.get("casos_etabs", {"cm": ["Dead", "SCP"], "cv": ["Live"], "sismo": ["SX", "SY"]})
    salida = {"fc_mpa": fc, "fy_mpa": fy,
              "combinaciones": e060.plan_combinaciones(casos["cm"], casos["cv"], casos.get("sismo"),
                                                       sismo_espectral=casos.get("sismo_espectral", True)),
              "vigas": [], "columnas": []}
    for v in cfg.get("vigas", []):
        fila = {"id": v["id"], "b_mm": v["b_mm"], "h_mm": v["h_mm"], "Mu_kNm": v.get("mu_knm"), "Vu_kN": v.get("vu_kn")}
        try:
            if v.get("mu_knm") is not None:
                fila["flexion"] = e060.flexion_viga(v["mu_knm"], v["b_mm"], v["h_mm"], fc, fy)
            if v.get("vu_kn") is not None:
                d = fila["flexion"]["d_mm"] if "flexion" in fila else v["h_mm"] - 60
                fila["cortante"] = e060.cortante_viga(v["vu_kn"], v["b_mm"], d, fc, fy, h_mm=v["h_mm"])
        except e060.ErrorE060 as exc:
            fila["error"] = str(exc)
        salida["vigas"].append(fila)
    for c in cfg.get("columnas", []):
        n, barra = c["barras"]
        as_total = n * e060.BARRAS[barra][1]
        fila = {"id": c["id"], "b_mm": c["b_mm"], "h_mm": c["h_mm"], "Pu_kN": c["pu_kn"],
                "refuerzo": f"{n} Ø {barra}", "As_mm2": as_total}
        try:
            fila["axial"] = e060.columna_axial(c["pu_kn"], c["b_mm"], c["h_mm"], as_total, fc, fy)
        except e060.ErrorE060 as exc:
            fila["error"] = str(exc)
        salida["columnas"].append(fila)
    return salida


# ------------------------------------------------------------------ ejecución --------------------------
def ejecutar(cfg: dict, confirmar: bool = False, boveda: Path | None = None, salida: Path | None = None) -> dict:
    """Calcula todo el proyecto. Sin confirmar: dry-run (no escribe nada)."""
    _requerir(cfg, "codigo", "nombre", "sismo")
    raiz = Path(boveda or ruta_boveda())
    base = nombre_archivo(f"{cfg['codigo']} {cfg['nombre']}")
    carpeta_salida = Path(salida or cfg.get("salida") or raiz / CARPETA_SALIDAS / nombre_archivo(cfg["codigo"]))
    pesos = pesos_por_nivel(cfg)
    sismo = analisis_sismico(cfg, pesos)
    concreto = diseno_concreto(cfg)
    metrado = None
    if cfg.get("modelo_ifc"):
        from . import metrados
        ruta_ifc = Path(cfg["modelo_ifc"]).expanduser()
        if not ruta_ifc.is_absolute() and cfg.get("_origen"):
            ruta_ifc = Path(cfg["_origen"]).parent / ruta_ifc
        metrado = metrados.metrar(str(ruta_ifc), por_piso=True)

    resultado = {"proyecto": base, "pesos": pesos, "sismo": sismo, "concreto": concreto, "metrado": metrado}
    pendientes = sorted({p for r in sismo.values() for p in r["estatica"]["pendientes"] + r["espectro"]["pendientes"]}
                        | {e030.PENDIENTES["derivas"], e030.PENDIENTES["desplazamientos"]})
    alertas = [a for r in sismo.values() for a in r["restricciones"] + r["estatica"]["notas"]]
    alertas += [f"{d}: método estático no aplicable (art. 33) → usa el modal espectral." for d, r in sismo.items()
                if not r["metodo_estatico_aplicable"]]
    alertas += [f"Viga {v['id']}: {v['error']}" for v in concreto["vigas"] if "error" in v]
    alertas += [f"Columna {c['id']}: Pu > φPn máx." for c in concreto["columnas"] if "axial" in c and not c["axial"]["cumple"]]
    alertas += [f"Columna {c['id']}: cuantía fuera de 1–6 %." for c in concreto["columnas"]
                if "axial" in c and not c["axial"]["cuantia_ok"]]
    resultado["pendientes"], resultado["alertas"] = pendientes, list(dict.fromkeys(alertas))

    archivos = {
        "excel": carpeta_salida / f"{nombre_archivo(cfg['codigo'])}_memoria_E030_E060.xlsx",
        "nota_proyecto": raiz / CARPETA_PROYECTOS / f"{base}.md",
        "nota_memoria": raiz / CARPETA_PROYECTOS / f"{nombre_archivo(cfg['codigo'])} - Memoria de calculo.md",
    }
    resultado["archivos"] = {k: str(v) for k, v in archivos.items()}
    resultado["resumen"] = _resumen(cfg, resultado)
    if not confirmar:
        resultado["dry_run"] = True
        resultado["mensaje"] = "Dry-run: nada escrito. Revisa el resumen y vuelve a ejecutar con confirmar=True."
        return resultado

    excel_rel = archivos["excel"]
    try:
        excel_rel = archivos["excel"].relative_to(raiz)
    except ValueError:
        pass
    memoria_a_excel(cfg, resultado, archivos["excel"])
    archivos["nota_memoria"].parent.mkdir(parents=True, exist_ok=True)
    archivos["nota_memoria"].write_text(nota_memoria(cfg, resultado, excel_rel), encoding="utf-8")
    if archivos["nota_proyecto"].exists():
        _agregar_bitacora(archivos["nota_proyecto"], resultado)
    else:
        archivos["nota_proyecto"].write_text(nota_proyecto(cfg, resultado, excel_rel), encoding="utf-8")
    bitacora.registrar("proyecto_ejecutar", f"{base}: E.020→E.030→E.060, memoria y notas", "escritura (bóveda/Excel)",
                       cfg.get("confirmado_por", "ingeniero (confirmar=true)"), resultado["resumen"]["linea"], boveda=raiz)
    resultado["dry_run"] = False
    resultado["mensaje"] = "Proyecto ejecutado: memoria Excel, nota de memoria y nota del proyecto escritas."
    return resultado


def _resumen(cfg: dict, r: dict) -> dict:
    lineas = []
    for d, s in r["sismo"].items():
        e = s["estatica"]
        lineas.append(f"{d}: R = {_g(e['R'])}, T = {_g(s['T_s'])} s, C = {_g(e['C'])}, V = {_num(e['V'], 1)} kN "
                      f"({_g(round(100 * e['V_sobre_P'], 2))} % P)")
    return {"P_total_kN": r["pesos"]["P_total_kN"], "hn_m": r["pesos"]["hn_m"],
            "parametros": {k: r["sismo"][next(iter(r["sismo"]))]["estatica"][k] for k in ("Z", "U", "S", "TP", "TL")},
            "direcciones": lineas, "alertas": len(r["alertas"]), "linea": " · ".join(lineas)}


# ------------------------------------------------------------------ Excel -------------------------------
NEGRITA, CABECERA = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")


def _tabla(hoja, titulos: list[str], filas: list[list]) -> None:
    hoja.append(titulos)
    for c in hoja[hoja.max_row]:
        c.font, c.fill = NEGRITA, CABECERA
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for f in filas:
        hoja.append(f)


def memoria_a_excel(cfg: dict, r: dict, destino: Path) -> str:
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    h = wb.active
    h.title = "Resumen"
    h.append([f"Memoria de cálculo (borrador JARVIS) — {r['proyecto']}"])
    h["A1"].font = Font(bold=True, size=13)
    h.append(["Debe ser revisada y firmada por el ingeniero responsable. Datos marcados 'supuesto' deben reemplazarse."])
    h.append([])
    for k in ("cliente", "ubicacion", "tipologia", "fase"):
        h.append([k.capitalize(), cfg.get(k, "")])
    h.append(["Normas", f"{e020.NORMA}; {e030.NORMA}; {e060.NORMA}"])
    h.append(["Generado", dt.datetime.now().strftime("%Y-%m-%d %H:%M")])
    h.append([])
    _tabla(h, ["Dirección", "Sistema", "R0", "Ia", "Ip", "R", "T (s)", "C", "C/R", "V (kN)", "V/P", "V mín. dinámica (kN)"],
           [[d, s["descripcion"], s["espectro"]["parametros"]["R0"], s["espectro"]["parametros"]["Ia"],
             s["espectro"]["parametros"]["Ip"], s["estatica"]["R"], s["T_s"], s["estatica"]["C"],
             s["estatica"]["C_sobre_R"], s["estatica"]["V"], s["estatica"]["V_sobre_P"], s["V_minima_dinamica_kN"]]
            for d, s in r["sismo"].items()])
    h.append([])
    h.append(["Alertas"])
    h[f"A{h.max_row}"].font = NEGRITA
    for a in r["alertas"] or ["Ninguna."]:
        h.append([a])
    h.append(["Pendientes normativos"])
    h[f"A{h.max_row}"].font = NEGRITA
    for p in r["pendientes"]:
        h.append([p])
    h.column_dimensions["A"].width, h.column_dimensions["B"].width = 22, 48

    hp = wb.create_sheet("Pesos E.020")
    hp.append([r["pesos"]["criterio"]])
    _tabla(hp, ["Nivel", "h acum. (m)", "Área (m²)", "CM (kPa)", "CV (kPa)", "CM (kN)", "CV (kN)", "% CV", "P (kN)", "Fuentes"],
           [[f["nivel"], f["altura_m"], f["area_m2"], f["cm_kpa"], f["cv_kpa"], f["CM_kN"], f["CV_kN"],
             f["fraccion_CV"], f["P_kN"], "; ".join(f"{k}: {v}" for k, v in f["fuentes"].items())]
            for f in r["pesos"]["filas"]])
    hp.append(["TOTAL", r["pesos"]["hn_m"], None, None, None, None, None, None, r["pesos"]["P_total_kN"]])
    hp[f"A{hp.max_row}"].font = NEGRITA

    for d, s in r["sismo"].items():
        e, p = s["estatica"], s["espectro"]["parametros"]
        hs = wb.create_sheet(f"Sismo {d}")
        hs.append([f"Dirección {d} — {s['descripcion']} — {e030.NORMA}"])
        hs["A1"].font = NEGRITA
        for k, v in [("Z", e["Z"]), ("U", e["U"]), ("Perfil", p["perfil"]), ("Vs30 (m/s)", p["vs30"]), ("S", e["S"]),
                     ("TP (s)", e["TP"]), ("TL (s)", e["TL"]), ("R0", p["R0"]), ("Ia", p["Ia"]), ("Ip", p["Ip"]),
                     ("R", e["R"]), ("T (s)", s["T_s"]), ("Fuente de T", s["fuente_T"]), ("C", e["C"]), ("C/R", e["C_sobre_R"]),
                     ("k", e["k"]), ("P (kN)", e["P"]), ("V = Z·U·C·S/R·P (kN)", e["V"]),
                     ("Método estático aplicable (art. 33)", "sí" if s["metodo_estatico_aplicable"] else "no")]:
            hs.append([k, v])
        hs.append([])
        _tabla(hs, ["Nivel", "h (m)", "P (kN)", "α", "F (kN)", "Cortante de piso (kN)"], [])
        acumulado, filas = 0.0, []
        for f, nv in zip(reversed(e["fuerzas_por_nivel"]), reversed(r["pesos"]["filas"])):
            acumulado += f["F"]
            filas.append([nv["nivel"], f["altura_m"], f["peso"], f["alfa"], f["F"], round(acumulado, 3)])
        for fila in filas:
            hs.append(fila)
        hs.column_dimensions["A"].width = 34

    he = wb.create_sheet("Espectros")
    direcciones = list(r["sismo"])
    base = r["sismo"][direcciones[0]]["espectro"]
    _tabla(he, ["T (s)"] + [f"Sa/g {d}" for d in direcciones], [])
    for i, t in enumerate(base["periodos_s"]):
        he.append([t] + [r["sismo"][d]["espectro"]["sa_g"][i] for d in direcciones])
    graf = ScatterChart()
    graf.title, graf.x_axis.title, graf.y_axis.title = "Espectro de diseño E.030-2026", "T (s)", "Sa/g"
    n = len(base["periodos_s"]) + 1
    for j in range(len(direcciones)):
        graf.series.append(Series(Reference(he, min_col=2 + j, min_row=1, max_row=n),
                                  Reference(he, min_col=1, min_row=2, max_row=n), title_from_data=True))
    he.add_chart(graf, "E2")

    c = r["concreto"]
    hc = wb.create_sheet("Combinaciones E.060")
    hc.append(["E.060 art. 9.2 — combinaciones para ETABS/SAP2000 (crear con csi_crear_combinaciones_e060)"])
    _tabla(hc, ["Nombre", "Tipo", "Casos / combinaciones"],
           [[p["nombre"], p["tipo"], ", ".join(f"{x}×{f:g}" for x, f in p["casos"]) if "casos" in p
             else ", ".join(p["combinaciones"])] for p in c["combinaciones"]])
    hc.column_dimensions["C"].width = 70

    hv = wb.create_sheet("Vigas E.060")
    hv.append([f"f'c = {c['fc_mpa']:g} MPa · fy = {c['fy_mpa']:g} MPa"])
    _tabla(hv, ["Viga", "b×h (mm)", "Mu (kN·m)", "d (mm)", "As req (mm²)", "As mín", "As máx", "As diseño",
                "Barras", "φMn (kN·m)", "Vu (kN)", "φVc (kN)", "Estribos", "Observación"], [])
    for v in c["vigas"]:
        fl, co = v.get("flexion", {}), v.get("cortante", {})
        hv.append([v["id"], f"{v['b_mm']}×{v['h_mm']}", v["Mu_kNm"], fl.get("d_mm"), fl.get("As_requerido_mm2"),
                   fl.get("As_min_mm2"), fl.get("As_max_mm2"), fl.get("As_diseno_mm2"),
                   fl["barras"][0]["barras"] if fl.get("barras") else "", fl.get("phi_Mn_kNm"), v["Vu_kN"],
                   co.get("phi_Vc_kN"), co.get("estribos"), v.get("error", "rige As mín." if fl.get("rige_minimo") else "")])
    hcol = wb.create_sheet("Columnas E.060")
    _tabla(hcol, ["Columna", "b×h (mm)", "Refuerzo", "As (mm²)", "Cuantía", "Pu (kN)", "φPn máx (kN)", "Uso %", "Cumple"],
           [[x["id"], f"{x['b_mm']}×{x['h_mm']}", x["refuerzo"], x["As_mm2"], x.get("axial", {}).get("cuantia"),
             x["Pu_kN"], x.get("axial", {}).get("phi_Pn_max_kN"), x.get("axial", {}).get("uso_%"),
             "sí" if x.get("axial", {}).get("cumple") else x.get("error", "no")] for x in c["columnas"]])
    if r.get("metrado"):
        hm = wb.create_sheet("Metrados IFC")
        hm.append([r["metrado"]["referencia"], r["metrado"]["archivo"]])
        _tabla(hm, ["Piso", "Clase IFC", "Partida", "Unidad", "Cantidad", "Elementos"],
               [[f["piso"], f["clase"], f["partida"], f["unidad"], f["cantidad"], f["elementos"]]
                for f in r["metrado"]["filas"]])
    wb.save(destino)
    return str(destino)


# ------------------------------------------------------------------ notas de la bóveda ------------------
def _tabla_md(titulos: list[str], filas: list[list]) -> str:
    salida = ["| " + " | ".join(titulos) + " |", "|" + "---|" * len(titulos)]
    salida += ["| " + " | ".join(_g(x) for x in f) + " |" for f in filas]
    return "\n".join(salida)


def nota_memoria(cfg: dict, r: dict, excel: Path) -> str:
    hoy = dt.date.today().isoformat()
    codigo = nombre_archivo(cfg["codigo"])
    s0 = next(iter(r["sismo"].values()))
    p = s0["espectro"]["parametros"]
    partes = [f"""---
tipo: documento
documento: memoria-calculo
proyecto: "{r['proyecto']}"
fecha: {hoy}
software: JARVIS (jarvis-bim) + ETABS
tags: [estructural, memoria, proyecto]
aliases: [{codigo} memoria]
actualizado: {hoy}
---
# Memoria de cálculo — {r['proyecto']}

> [!warning] Borrador generado por JARVIS
> Se regenera con `proyecto_ejecutar` (no edites a mano). Debe ser revisado y firmado por el ingeniero responsable.
> Los datos marcados *supuesto* deben reemplazarse por los del proyecto o del modelo ETABS.
> Excel: `{Path(excel).as_posix()}`

Proyecto: [[{r['proyecto']}]] · Plantilla: [[Plantilla - Memoria de calculo]]

## 1. Descripción
- Cliente: {cfg.get('cliente', '—')} · Ubicación: {cfg.get('ubicacion', '—')} · Tipología: {cfg.get('tipologia', '—')}
- {r['pesos']['pisos']} niveles · hn = {_num(r['pesos']['hn_m'])} m · f'c = {r['concreto']['fc_mpa']:g} MPa · fy = {r['concreto']['fy_mpa']:g} MPa

## 2. Normas
[[E.020 - Cargas]] · [[E.030 - Diseno Sismorresistente]] ({e030.NORMA}) · [[E.060 - Concreto Armado]] ({e060.NORMA})

## 3. Pesos por nivel (E.020 + E.030 art. 31)
{r['pesos']['criterio']}.

""",
        _tabla_md(["Nivel", "h (m)", "Área (m²)", "CM (kPa)", "CV (kPa)", "CM (kN)", "CV (kN)", "%CV", "P (kN)"],
                  [[f["nivel"], _num(f["altura_m"]), f["area_m2"], f["cm_kpa"], f["cv_kpa"], _num(f["CM_kN"], 1),
                    _num(f["CV_kN"], 1), f"{int(100 * f['fraccion_CV'])} %", _num(f["P_kN"], 1)]
                   for f in r["pesos"]["filas"]]),
        f"\n\n**P total = {_num(r['pesos']['P_total_kN'], 1)} kN**\n\nFuentes: " + "; ".join(sorted(
            {f"{k}: {v}" for f in r["pesos"]["filas"] for k, v in f["fuentes"].items()})) + ".\n",
        f"""
## 4. Parámetros sísmicos (E.030-2026)
Z = {_g(p['Z'])} · U = {_g(p['U'])} · perfil {p['perfil']} (Vs30 = {_g(p['vs30'])} m/s) · S = {_g(p['S'])} · TP = {_g(p['TP'])} s · TL = {_g(p['TL'])} s · categoría {p['categoria']}

""",
        _tabla_md(["Dir.", "Sistema", "R0", "Ia", "Ip", "R", "T (s)", "C", "C/R", "V (kN)", "V/P", "V mín. din. (kN)"],
                  [[d, s["descripcion"], s["espectro"]["parametros"]["R0"], s["espectro"]["parametros"]["Ia"],
                    s["espectro"]["parametros"]["Ip"], s["estatica"]["R"], s["T_s"], s["estatica"]["C"],
                    s["estatica"]["C_sobre_R"], _num(s["estatica"]["V"], 1), f"{_g(round(100 * s['estatica']['V_sobre_P'], 2))} %",
                    _num(s["V_minima_dinamica_kN"], 1)] for d, s in r["sismo"].items()]),
        "\n\nFórmula: V = Z·U·C·S/R·P con C/R ≥ 0,11 (arts. 34–35); T = hn/CT (art. 36); "
        "cortante dinámica mínima 80 % (regular) / 90 % (irregular) de la estática (art. 44).\n",
    ]
    for d, s in r["sismo"].items():
        e = s["estatica"]
        partes.append(f"\n### Dirección {d}: fuerzas por nivel (k = {_g(e['k'])})\n")
        partes.append(_tabla_md(["Nivel", "h (m)", "P (kN)", "α", "F (kN)"],
                                [[nv["nivel"], f["altura_m"], _num(f["peso"], 1), f["alfa"], _num(f["F"], 1)]
                                 for f, nv in zip(e["fuerzas_por_nivel"], r["pesos"]["filas"])]))
        partes.append(f"\n\n- {s['fuente_T']} · método estático {'aplicable' if s['metodo_estatico_aplicable'] else '**no aplicable**'} (art. 33)\n")
        if s.get("escalamiento"):
            partes.append(f"- Escalamiento dinámico: factor {s['escalamiento']['factor_escala']}\n")
        if s.get("derivas"):
            dv = s["derivas"]
            partes.append(f"- Derivas ({dv['factor']}, límite {dv['limite']}): {'cumple' if dv['cumple_todo'] else '**NO cumple**'}\n")
    c = r["concreto"]
    partes.append("\n## 5. Combinaciones de carga (E.060 art. 9.2)\n")
    partes.append(_tabla_md(["Nombre", "Tipo", "Casos"],
                            [[x["nombre"], x["tipo"], ", ".join(f"{n}×{f:g}" for n, f in x["casos"]) if "casos" in x
                              else ", ".join(x["combinaciones"])] for x in c["combinaciones"]]))
    if c["vigas"]:
        partes.append("\n\n## 6. Vigas (E.060 arts. 10 y 11)\n")
        partes.append(_tabla_md(["Viga", "b×h", "Mu (kN·m)", "As diseño (mm²)", "Barras", "φMn", "Vu (kN)", "Estribos", "Obs."],
                                [[v["id"], f"{v['b_mm']}×{v['h_mm']}", v["Mu_kNm"],
                                  v.get("flexion", {}).get("As_diseno_mm2"),
                                  v["flexion"]["barras"][0]["barras"] if v.get("flexion", {}).get("barras") else "",
                                  v.get("flexion", {}).get("phi_Mn_kNm"), v["Vu_kN"], v.get("cortante", {}).get("estribos"),
                                  v.get("error", "rige As mín." if v.get("flexion", {}).get("rige_minimo") else "")]
                                 for v in c["vigas"]]))
        partes.append("\n\nAs mín = 0,22·√f'c/fy·b·d (10.5.2); As máx = 0,75·Asb (10.3.4); Vc = 0,17·√f'c·b·d (11.3.1.1). "
                      "Confinamiento y diseño por capacidad del cap. 21 aparte.\n")
    if c["columnas"]:
        partes.append("\n## 7. Columnas — carga axial máxima (E.060 10.3.6 y 10.9.1)\n")
        partes.append(_tabla_md(["Columna", "b×h", "Refuerzo", "Cuantía", "Pu (kN)", "φPn máx (kN)", "Uso", "Cumple"],
                                [[x["id"], f"{x['b_mm']}×{x['h_mm']}", x["refuerzo"], x.get("axial", {}).get("cuantia"),
                                  x["Pu_kN"], x.get("axial", {}).get("phi_Pn_max_kN"), f"{_g(x.get('axial', {}).get('uso_%'))} %",
                                  "sí" if x.get("axial", {}).get("cumple") else "**no**"] for x in c["columnas"]]))
        partes.append("\n\nSolo axial: con flexión verifica con el diagrama de interacción en ETABS.\n")
    partes.append("\n## 8. Alertas\n" + ("\n".join(f"- {a}" for a in r["alertas"]) or "- Ninguna.") + "\n")
    partes.append("\n## 9. Pendientes normativos\n" + "\n".join(f"- [!todo] {p}" for p in r["pendientes"]) + "\n")
    partes.append(f"""
## 10. Siguientes pasos en ETABS
1. `csi_crear_espectro_e030` (dry-run → confirmar) con los parámetros de la sección 4, una función por dirección.
2. `csi_crear_combinaciones_e060` con los casos del modelo (sección 5).
3. `csi_correr_analisis(confirmar=True)` → `csi_reacciones_base` para V dinámica y `csi_derivas` para distorsiones.
4. Copia V dinámica y derivas en `resultados_etabs` del JSON del proyecto y vuelve a ejecutar.

↑ [[{r['proyecto']}]] · [[Indice de proyectos]] · [[JARVIS - Flujos de trabajo]]
""")
    return "".join(partes)


def nota_proyecto(cfg: dict, r: dict, excel: Path) -> str:
    hoy = dt.date.today().isoformat()
    codigo = nombre_archivo(cfg["codigo"])
    s = cfg["sismo"]
    s0 = next(iter(r["sismo"].values()))
    p = s0["espectro"]["parametros"]
    sistemas = " / ".join(f"{d}: {x['sistema']}" for d, x in r["sismo"].items())
    return f"""---
tipo: proyecto
estado: activo
cliente: "{cfg.get('cliente', '')}"
codigo_proyecto: "{cfg['codigo']}"
fase: {cfg.get('fase', 'expediente')}
tipologia: "{cfg.get('tipologia', '')}"
ubicacion: "{cfg.get('ubicacion', '')}"
zona_sismica: {s['zona']}
perfil_suelo: {p['perfil']}
vs30: {p['vs30'] if p['vs30'] is not None else ''}
categoria_uso: {p['categoria']}
sistema_estructural: "{sistemas}"
fecha_inicio: {hoy}
fecha_entrega: {cfg.get('fecha_entrega', '')}
cde: "{cfg.get('cde', '')}"
ppc:
cpi:
spi:
choques_abiertos:
rfis_abiertos:
tags: [proyecto]
aliases: [{codigo}]
actualizado: {hoy}
---
# {r['proyecto']}

## 1. Datos generales
- Cliente / entidad: {cfg.get('cliente', '—')}
- Ubicación: {cfg.get('ubicacion', '—')} · Tipología: {cfg.get('tipologia', '—')}
- Tipo de inversión: {cfg.get('inversion', '—')} → [[Invierte.pe - Ciclo de inversion]] · ¿BIM obligatorio? → [[Plan BIM Peru]]
- Datos de entrada: `{Path(cfg.get('_origen', 'config.json')).name}`

## 2. Documentos de información
- [ ] EIR recibido/elaborado → [[Plantilla - EIR]]
- [ ] PEB → [[Plantilla - BEP]]
- [ ] MIDP/TIDP → [[MIDP y TIDP - Planes de entrega]]
- [ ] CDE configurado → [[CDE - Entorno Comun de Datos]]

## 3. Modelos (federación)
| Disciplina | Archivo | Software | Estado | Revisión |
|---|---|---|---|---|
| ARQ | | [[Revit]] | S0 | P01 |
| EST | {Path(cfg['modelo_ifc']).name if cfg.get('modelo_ifc') else ''} | [[Revit]] / [[ETABS]] | S0 | P01 |
| MEP | | [[Revit]] | S0 | P01 |

## 4. Parámetros estructurales
Z = {_g(p['Z'])} · U = {_g(p['U'])} · S = {_g(p['S'])} · TP = {_g(p['TP'])} s · TL = {_g(p['TL'])} s · {sistemas} → [[E.030 - Diseno Sismorresistente]]

Resultado JARVIS: {r['resumen']['linea']} → [[{codigo} - Memoria de calculo]] · Excel `{Path(excel).as_posix()}`

## 5. Hitos
| Hito | Fecha | Entregables |
|---|---|---|

## 6. Riesgos principales → [[Gestion de riesgos]]
{chr(10).join(f'- {a}' for a in r['alertas']) or '- (sin alertas normativas)'}

## 7. Bitácora
- {hoy}: creación del proyecto y primera ejecución de JARVIS (E.020 → E.030-2026 → E.060).

↑ [[Indice de proyectos]]
"""


def _agregar_bitacora(nota: Path, r: dict) -> None:
    texto = nota.read_text(encoding="utf-8")
    linea = f"- {dt.date.today().isoformat()}: nueva ejecución de JARVIS — {r['resumen']['linea']}"
    marcador = "\n↑ [[Indice de proyectos]]"
    texto = texto.replace(marcador, f"{linea}\n{marcador}", 1) if marcador in texto else texto + f"\n{linea}\n"
    nota.write_text(texto, encoding="utf-8")
