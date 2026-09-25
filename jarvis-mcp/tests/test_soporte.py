from pathlib import Path

import pytest

from jarvis_bim import bitacora, e030, nomenclatura, reportes
from jarvis_bim.boveda import Boveda, ruta_boveda


@pytest.fixture(scope="module")
def boveda():
    return Boveda(ruta_boveda())


def test_boveda_carga_notas(boveda):
    est = boveda.estadisticas()
    assert est["notas"] > 250
    assert "30 NORMAS" in est["capas"]


def test_busqueda_encuentra_e030(boveda):
    res = boveda.buscar("periodo del suelo Ts zona 4 sismorresistente")
    assert any(r["nota"].startswith("E.030") for r in res[:3])
    assert boveda.buscar("norma sismorresistente derivas")[0]["nota"].startswith("E.030")


def test_busqueda_sin_tildes(boveda):
    assert boveda.buscar("albanileria densidad de muros")[0]["nota"] == "E.070 - Albanileria"


def test_leer_nota_con_retroenlaces(boveda):
    n = boveda.leer("E.030 - Diseno Sismorresistente")
    assert "RM 183-2026" in n["contenido"]
    assert "MOC Normas y Estandares" in n["retroenlaces"]


def test_leer_nota_inexistente_sugiere(boveda):
    with pytest.raises(KeyError):
        boveda.leer("E.030 sismo inexistente")


def test_nomenclatura_valida():
    r = nomenclatura.validar_nombre("PRJ01-ABC-ZZ-01-M3-S-0001.rvt")
    assert r["valido"], r["errores"]
    assert r["descripcion"]["rol"] == "estructuras"


def test_nomenclatura_invalida():
    r = nomenclatura.validar_nombre("modelo final v2.rvt")
    assert not r["valido"]
    r2 = nomenclatura.validar_nombre("PRJ01-ABC-ZZ-01-XX-Q9-12")
    assert not r2["valido"] and len(r2["errores"]) >= 2


def test_bitacora(tmp_path):
    ruta = bitacora.registrar("prueba", "acción de prueba", boveda=tmp_path)
    bitacora.registrar("prueba", "segunda acción", modo="escritura", confirmado_por="Jhordan", boveda=tmp_path)
    texto = ruta.read_text(encoding="utf-8")
    assert texto.count("| prueba |") == 2
    assert texto.index("segunda acción") < texto.index("↑ [[Indice del Diario JARVIS]]")


def test_espectro_a_excel(tmp_path):
    from openpyxl import load_workbook
    esp = e030.espectro(4, "C", "concreto_dual", perfil="S2", vs30=450)
    ruta = reportes.espectro_a_excel(esp, str(tmp_path / "espectro"))
    wb = load_workbook(ruta)
    assert wb["Espectro"].max_row == len(esp["periodos_s"]) + 1


def _modelo_ifc_minimo(ruta: Path):
    ifcopenshell = pytest.importorskip("ifcopenshell")
    import ifcopenshell.api as api
    m = api.run("project.create_file", version="IFC4")
    proyecto = api.run("root.create_entity", m, ifc_class="IfcProject", name="Prueba JARVIS")
    api.run("unit.assign_unit", m)
    sitio = api.run("root.create_entity", m, ifc_class="IfcSite", name="Sitio")
    edificio = api.run("root.create_entity", m, ifc_class="IfcBuilding", name="Edificio")
    piso = api.run("root.create_entity", m, ifc_class="IfcBuildingStorey", name="Piso 1")
    api.run("aggregate.assign_object", m, relating_object=proyecto, products=[sitio])
    api.run("aggregate.assign_object", m, relating_object=sitio, products=[edificio])
    api.run("aggregate.assign_object", m, relating_object=edificio, products=[piso])
    muro = api.run("root.create_entity", m, ifc_class="IfcWall", name="Muro 1")
    api.run("spatial.assign_container", m, relating_structure=piso, products=[muro])
    pset = api.run("pset.add_pset", m, product=muro, name="Pset_WallCommon")
    api.run("pset.edit_pset", m, pset=pset, properties={"LoadBearing": True, "FireRating": "F60"})
    m.write(str(ruta))
    return muro.GlobalId


def test_ifc_resumen_y_propiedades(tmp_path):
    from jarvis_bim import ifc
    ruta = tmp_path / "prueba.ifc"
    guid = _modelo_ifc_minimo(ruta)
    r = ifc.resumen(str(ruta))
    assert r["esquema"] == "IFC4" and r["clases"].get("IfcWall") == 1
    e = ifc.elementos(str(ruta), "IfcWall")
    assert e["total"] == 1 and e["elementos"][0]["contenedor"] == "Piso 1"
    p = ifc.propiedades(str(ruta), guid)
    assert p["psets"]["Pset_WallCommon"]["FireRating"] == "F60"


def test_busqueda_devuelve_seccion_del_texto_oficial(boveda):
    if not (boveda.raiz / "30 NORMAS" / "Textos oficiales" / "Texto oficial - E060.md").exists():
        pytest.skip("texto oficial de la E.060 no disponible")
    res = boveda.buscar("area minima de refuerzo por traccion secciones rectangulares", carpeta="30 NORMAS/Textos")
    assert res[0]["nota"] == "Texto oficial - E060"
    assert res[0]["seccion"].startswith("10.5.2")
