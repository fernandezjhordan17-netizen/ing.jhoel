import math

import pytest

ifcopenshell = pytest.importorskip("ifcopenshell")
import ifcopenshell.api as api
import ifcopenshell.util.unit as uu

from jarvis_bim import metrados


def _modelo(ruta):
    m = api.run("project.create_file", version="IFC4")
    proyecto = api.run("root.create_entity", m, ifc_class="IfcProject", name="Metrado")
    api.run("unit.assign_unit", m)
    esc = uu.calculate_unit_scale(m)
    piso = api.run("root.create_entity", m, ifc_class="IfcBuildingStorey", name="Piso 1")
    api.run("aggregate.assign_object", m, relating_object=proyecto, products=[piso])
    for i, vol in enumerate((0.36, 0.36)):
        col = api.run("root.create_entity", m, ifc_class="IfcColumn", name=f"C{i}")
        api.run("spatial.assign_container", m, relating_structure=piso, products=[col])
        q = api.run("pset.add_qto", m, product=col, name="Qto_ColumnBaseQuantities")
        api.run("pset.edit_qto", m, qto=q, properties={"NetVolume": vol})
    viga = api.run("root.create_entity", m, ifc_class="IfcBeam", name="V1")  # sin Qto
    api.run("spatial.assign_container", m, relating_structure=piso, products=[viga])
    barra = api.run("root.create_entity", m, ifc_class="IfcReinforcingBar", name="Ø16")
    barra.NominalDiameter = 0.016 / esc
    barra.BarLength = 6.0 / esc
    api.run("spatial.assign_container", m, relating_structure=piso, products=[barra])
    puerta = api.run("root.create_entity", m, ifc_class="IfcDoor", name="P1")
    api.run("spatial.assign_container", m, relating_structure=piso, products=[puerta])
    m.write(str(ruta))


def test_metrado(tmp_path):
    ruta = tmp_path / "m.ifc"
    _modelo(ruta)
    r = metrados.metrar(str(ruta))
    filas = {f["clase"]: f for f in r["filas"]}
    assert math.isclose(filas["IfcColumn"]["cantidad"], 0.72) and filas["IfcColumn"]["unidad"] == "m³"
    assert math.isclose(filas["IfcReinforcingBar"]["cantidad"], math.pi / 4 * 0.016 ** 2 * 6 * 7850, rel_tol=1e-3)
    assert filas["IfcDoor"]["cantidad"] == 1 and filas["IfcDoor"]["unidad"] == "und"
    assert filas["IfcBeam"]["sin_dato"] == 1 and any("IfcBeam" in a for a in r["avisos"])


def test_metrado_por_piso_y_excel(tmp_path):
    from openpyxl import load_workbook
    ruta = tmp_path / "m.ifc"
    _modelo(ruta)
    r = metrados.metrar(str(ruta), por_piso=True)
    assert all(f["piso"] == "Piso 1" for f in r["filas"])
    x = metrados.metrado_a_excel(r, str(tmp_path / "metrado"))
    assert load_workbook(x)["Metrado"].max_row == len(r["filas"]) + 1


def test_duplicado_con_error_de_unidades(tmp_path):
    m = api.run("project.create_file", version="IFC4")
    proyecto = api.run("root.create_entity", m, ifc_class="IfcProject", name="Dup")
    api.run("unit.assign_unit", m)
    muro = api.run("root.create_entity", m, ifc_class="IfcWall", name="M1")
    q = api.run("pset.add_qto", m, product=muro, name="Qto_WallBaseQuantities")
    area_mal = m.create_entity("IfcQuantityArea", Name="NetSideArea", AreaValue=45252070.0)
    area_bien = m.create_entity("IfcQuantityArea", Name="NetSideArea", AreaValue=45.25207)
    q.Quantities = (area_mal, area_bien)
    ruta = tmp_path / "dup.ifc"
    m.write(str(ruta))
    r = metrados.metrar(str(ruta))
    assert math.isclose(r["filas"][0]["cantidad"], 45.252, rel_tol=1e-4)
    assert any("error de unidades" in a for a in r["avisos"])


def test_valor_unico_implausible(tmp_path):
    m = api.run("project.create_file", version="IFC4")
    api.run("root.create_entity", m, ifc_class="IfcProject", name="X")
    api.run("unit.assign_unit", m)
    muro = api.run("root.create_entity", m, ifc_class="IfcWall", name="M2")
    q = api.run("pset.add_qto", m, product=muro, name="Qto_WallBaseQuantities")
    q.Quantities = (m.create_entity("IfcQuantityArea", Name="NetSideArea", AreaValue=29292000.0),)
    ruta = tmp_path / "x.ifc"
    m.write(str(ruta))
    assert math.isclose(metrados.metrar(str(ruta))["filas"][0]["cantidad"], 29.292, rel_tol=1e-4)
