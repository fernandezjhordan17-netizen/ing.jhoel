import json
import math
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from jarvis_bim import e020, proyecto
from jarvis_bim.boveda import Boveda, ruta_boveda
from jarvis_bim.proyecto import ErrorProyecto

EJEMPLO = Path(__file__).resolve().parents[2] / "jarvis" / "proyectos" / "PRY001_ejemplo.json"


@pytest.fixture
def boveda(tmp_path):
    copia = tmp_path / "JARVIS-BIM"
    shutil.copytree(ruta_boveda(), copia, ignore=shutil.ignore_patterns(".obsidian"))
    return copia


@pytest.fixture
def cfg():
    datos = proyecto.cargar(EJEMPLO)
    datos["codigo"] = "PRYTEST"  # la bóveda real ya contiene PRY001
    return datos


def test_e020_tablas():
    assert e020.peso_aligerado(0.20) == 3.0 and e020.PESO_CONCRETO_ARMADO == 24.0
    assert e020.carga_viva("vivienda") == 2.0 and e020.carga_viva("techo") == 1.0
    with pytest.raises(e020.ErrorE020):
        e020.peso_aligerado(0.22)


def test_pesos_por_nivel(cfg):
    p = proyecto.pesos_por_nivel(cfg)
    piso = p["filas"][0]
    # 3,0 (aligerado 0,20 m) + 1,0 + 1,5 + 2,5 = 8,0 kPa; CV vivienda 2,0 kPa; categoría C → 25 %
    assert piso["cm_kpa"] == 8.0 and piso["cv_kpa"] == 2.0
    assert piso["P_kN"] == 280 * 8.0 + 0.25 * 280 * 2.0
    azotea = p["filas"][-1]
    assert azotea["cv_kpa"] == 1.0 and azotea["cm_kpa"] == 6.5 and azotea["fraccion_CV"] == 0.25
    assert p["hn_m"] == 13.5 and p["P_total_kN"] == 4 * 2380 + 1890


def test_cortante_por_direccion_a_mano(cfg):
    r = proyecto.ejecutar(cfg)
    assert r["dry_run"]
    for d, R in (("X", 7.0), ("Y", 6.0)):
        e = r["sismo"][d]["estatica"]
        T = 13.5 / 60
        assert math.isclose(r["sismo"][d]["T_s"], round(T, 4))
        esperado = 0.45 * 1.0 * 2.5 * 1.065 / R * 11410
        assert math.isclose(e["V"], round(esperado, 3), abs_tol=1e-3)
        assert math.isclose(sum(f["F"] for f in e["fuerzas_por_nivel"]), e["V"], abs_tol=0.01)
        assert r["sismo"][d]["metodo_estatico_aplicable"]
        assert r["sismo"][d]["V_minima_dinamica_kN"] == round(0.8 * e["V"], 2)


def test_dry_run_no_escribe(cfg, boveda):
    r = proyecto.ejecutar(cfg, boveda=boveda)
    assert all(not Path(a).exists() for a in r["archivos"].values())


def test_ejecucion_confirmada_escribe_y_enlaza(cfg, boveda):
    r = proyecto.ejecutar(cfg, confirmar=True, boveda=boveda)
    assert not r["dry_run"]
    for a in r["archivos"].values():
        assert Path(a).exists(), a
    wb = load_workbook(r["archivos"]["excel"])
    assert {"Resumen", "Pesos E.020", "Sismo X", "Sismo Y", "Espectros", "Vigas E.060", "Columnas E.060"} <= set(wb.sheetnames)
    nota = Path(r["archivos"]["nota_proyecto"]).read_text(encoding="utf-8")
    assert "[[PRYTEST - Memoria de calculo]]" in nota and "zona_sismica: 4" in nota
    memoria = Path(r["archivos"]["nota_memoria"]).read_text(encoding="utf-8")
    assert "[!todo]" in memoria and "Borrador generado por JARVIS" in memoria
    # las notas nuevas quedan conectadas en la red de la bóveda
    b = Boveda(boveda)
    assert r["proyecto"] in b.leer("PRYTEST - Memoria de calculo")["retroenlaces"]
    assert list((boveda / "00 INICIO" / "Diario JARVIS").glob("20*.md"))
    # segunda ejecución: no pisa la nota del proyecto, agrega una línea de bitácora
    proyecto.ejecutar(cfg, confirmar=True, boveda=boveda)
    nota2 = Path(r["archivos"]["nota_proyecto"]).read_text(encoding="utf-8")
    assert nota2.count("nueva ejecución de JARVIS") == 1


def test_resultados_etabs_escalamiento_y_derivas(cfg):
    cfg["resultados_etabs"] = {"X": {"v_dinamica_kN": 1200.0, "derivas_elasticas": [0.0008, 0.0011, 0.0012, 0.0011, 0.0009]}}
    r = proyecto.ejecutar(cfg)
    esc = r["sismo"]["X"]["escalamiento"]
    assert math.isclose(esc["factor_escala"], round(0.8 * r["sismo"]["X"]["estatica"]["V"] / 1200.0, 4))
    assert r["sismo"]["X"]["derivas"]["filas"][2]["deriva_inelastica"] == round(0.75 * 7 * 0.0012, 6)


def test_alertas_y_errores(cfg):
    cfg["sismo"]["direcciones"]["X"]["irregularidades"] = ["torsional_extrema"]
    cfg["columnas"][0]["pu_kn"] = 5000
    r = proyecto.ejecutar(cfg)
    assert any("Tabla 13" in a for a in r["alertas"])
    assert any("C-1" in a for a in r["alertas"])
    del cfg["niveles"][0]["altura_entrepiso_m"]
    with pytest.raises(ErrorProyecto):
        proyecto.ejecutar(cfg)


def test_nombre_archivo_sin_tildes():
    assert proyecto.nombre_archivo("PRY002: Colegio Chiclayo/Etapa ñ") == "PRY002- Colegio Chiclayo-Etapa n"
